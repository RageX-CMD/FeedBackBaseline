# -*- coding: utf-8 -*-
"""Erstellt das docx-Dokument fuer eine Eventwoche."""

import os
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import OUTPUT_DIR


def build_docx(week_label, week_number, week_start, week_end, events):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc = Document()

    title = doc.add_heading(f"{week_label} - Feedback-Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta = doc.add_paragraph()
    meta.add_run("Zeitraum: ").bold = True
    meta.add_run(f"{week_start.strftime('%d.%m.%Y')} - {week_end.strftime('%d.%m.%Y')}")
    meta.add_run("\nErstellt am: ").bold = True
    meta.add_run(date.today().strftime("%d.%m.%Y"))

    doc.add_paragraph("")

    for ev in events:
        doc.add_heading(f"{ev['event_order']} - {ev['name']}", level=2)
        p = doc.add_paragraph()
        feedback_text = (ev["feedback"] or "").strip()
        if feedback_text:
            p.add_run(feedback_text)
        else:
            run = p.add_run("(kein Feedback vorhanden)")
            run.italic = True
        doc.add_paragraph("")

    filename = f"Eventwoche_{week_number}_{week_start.isoformat()}.docx"
    path = os.path.join(OUTPUT_DIR, filename)
    doc.save(path)
    return path
